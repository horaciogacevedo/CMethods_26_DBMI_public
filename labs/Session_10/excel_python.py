"""Create a small Excel workbook with openpyxl."""
from openpyxl import Workbook

if __name__ == '__main__':
    print("My Excel Example")
    workbook = Workbook()
    ws = workbook.active
    ws.title = 'worksheet_1'
    ws.sheet_properties.tabColor = '1072BA'
    ws['A1'] = 42
    ws['A2'] = 12
    ws['A3'] = '=SUM(A1,A2)'
    ws2 = workbook.create_sheet(title='worksheet_2')
    ws2['A1'] = 3.42
    ws2.append([1, 2, 3])
    ws2.cell(column=2, row=1, value=15)
    workbook.save('sample.xlsx')
    print("Finishing off")
